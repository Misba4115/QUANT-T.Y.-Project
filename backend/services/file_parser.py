"""File-type parsing and financial-row normalization."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from schemas.schemas import (
    CryptoCreate,
    EquityCreate,
    ETFCreate,
    MutualFundCreate,
    AssetTransactionCreate,
)
from services.cas_parser import parse_cas_pdf
from services.csv_parser import header_mapping, normalize_header, parse_csv_document, parse_csv_rows
from services.statement_parser import normalize_transaction_row, parse_date

PARSER_VERSION = "2026.08.7"

TARGET_REQUIREMENTS: dict[str, list[str]] = {
    "equity_holdings": ["symbol", "quantity", "avg_price"],
    "mutual_fund_holdings": ["mf_name", "units", "avg_nav"],
    "etf_holdings": ["symbol", "etf_name", "quantity", "avg_price"],
    "crypto_holdings": ["symbol", "crypto_name", "quantity", "avg_price"],
    "bank_transactions": ["tx_date", "description"],
    "mixed_asset_transactions": ["asset_type", "tx_date", "tx_type", "quantity", "price"],
    "equity_transactions": ["tx_date", "tx_type", "symbol", "quantity", "price"],
    "mf_transactions": ["tx_date", "tx_type", "mf_name", "units", "nav"],
    "etf_transactions": ["tx_date", "tx_type", "symbol", "quantity", "price"],
    "crypto_transactions": ["tx_date", "tx_type", "symbol", "quantity", "price"],
}


@dataclass(frozen=True)
class ParsedFinancialFile:
    target: str
    rows: list[dict[str, Any]]
    source_to_canonical: dict[str, str]
    recognized_columns: list[str]
    ignored_columns: list[str]
    mapping_method: str = "deterministic"


def parse_xlsx_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    rows, _mapping = parse_xlsx_document(file_bytes)
    return rows


def parse_xlsx_document(
    file_bytes: bytes,
    overrides: dict[str, str] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("The XLSX workbook could not be read") from exc
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        source_headers = list(next(values))
        mapping = header_mapping(source_headers, overrides)
        headers = [mapping.get(str(value or "").strip(), normalize_header(value)) for value in source_headers]
    except StopIteration as exc:
        raise ValueError("The XLSX workbook is empty") from exc
    rows = []
    for values_row in values:
        row = {headers[index]: value for index, value in enumerate(values_row) if index < len(headers)}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows, mapping


def detect_target(headers: set[str]) -> str:
    if (
        {"asset_type", "tx_date", "tx_type"}.issubset(headers)
        and {"quantity", "units"} & headers
        and {"price", "nav"} & headers
    ):
        return "mixed_asset_transactions"
    if "mf_name" in headers or {"scheme_code", "units", "avg_nav"}.issubset(headers):
        return "mutual_fund_holdings"
    if "etf_name" in headers:
        return "etf_holdings"
    if "crypto_name" in headers:
        return "crypto_holdings"
    if "equity_name" in headers or {"symbol", "quantity", "avg_price"}.issubset(headers):
        return "equity_holdings"
    if {"tx_date", "description"}.issubset(headers) and ({"amount"}.issubset(headers) or "debit" in headers or "credit" in headers):
        return "bank_transactions"
    raise ValueError(
        "The file structure is not recognized. File uploads are for bank transactions and portfolio "
        "holdings (equities, mutual funds, ETFs, and crypto). Bank accounts and financial goals are "
        "entered directly in Input Data."
    )


def _number(value: Any) -> Any:
    if isinstance(value, str):
        return value.replace(",", "").replace("₹", "").strip()
    return value


def normalize_row(target: str, row: dict[str, Any]) -> dict[str, Any]:
    cleaned = {key: _number(value) for key, value in row.items()}
    if target == "bank_transactions":
        return normalize_transaction_row(cleaned)
    if target == "mixed_asset_transactions":
        return _normalize_mixed_asset_transaction(cleaned)
    if target in {"equity_transactions", "mf_transactions", "etf_transactions", "crypto_transactions"}:
        asset_names = {
            "equity_transactions": "Equity",
            "mf_transactions": "MutualFund",
            "etf_transactions": "ETF",
            "crypto_transactions": "Crypto",
        }
        adapted = {
            **cleaned,
            "asset_type": asset_names[target],
            "asset_name": cleaned.get("asset_name") or cleaned.get("mf_name"),
            "quantity": cleaned.get("quantity") or cleaned.get("units"),
            "price": cleaned.get("price") or cleaned.get("nav"),
        }
        return _normalize_mixed_asset_transaction(adapted)
    models = {
        "equity_holdings": EquityCreate,
        "mutual_fund_holdings": MutualFundCreate,
        "etf_holdings": ETFCreate,
        "crypto_holdings": CryptoCreate,
    }
    model = models[target]
    allowed = set(model.model_fields)
    payload = {key: value for key, value in cleaned.items() if key in allowed and value not in ("", None)}
    if target == "equity_holdings" and payload.get("symbol") and not payload.get("equity_name"):
        payload["equity_name"] = cleaned.get("asset_name") or payload["symbol"]
    if target == "etf_holdings" and payload.get("symbol") and not payload.get("etf_name"):
        payload["etf_name"] = cleaned.get("asset_name") or payload["symbol"]
    if target == "crypto_holdings" and payload.get("symbol") and not payload.get("crypto_name"):
        payload["crypto_name"] = cleaned.get("asset_name") or payload["symbol"]
    normalized = model.model_validate(payload).model_dump(mode="json")
    price_key = "current_nav" if target == "mutual_fund_holdings" else "current_price"
    if cleaned.get(price_key) not in (None, ""):
        normalized[price_key] = float(cleaned[price_key])
        normalized["price_source"] = "imported_statement"
    return normalized


def _normalize_mixed_asset_transaction(row: dict[str, Any]) -> dict[str, Any]:
    asset_key = "".join(character for character in str(row.get("asset_type") or "").lower() if character.isalnum())
    target_tables = {
        "stock": "equity_transactions",
        "equity": "equity_transactions",
        "share": "equity_transactions",
        "mutualfund": "mf_transactions",
        "mf": "mf_transactions",
        "etf": "etf_transactions",
        "crypto": "crypto_transactions",
        "cryptocurrency": "crypto_transactions",
    }
    target_table = target_tables.get(asset_key)
    if not target_table:
        raise ValueError(f"Unsupported asset type '{row.get('asset_type')}'")

    raw_type = "".join(character for character in str(row.get("tx_type") or "").lower() if character.isalnum())
    if any(token in raw_type for token in {"buy", "purchase", "sip", "invest", "reinvest", "contribution"}):
        tx_type = "buy"
    elif any(token in raw_type for token in {"sell", "sale", "redeem", "redemption", "withdraw", "liquidat"}):
        tx_type = "sell"
    else:
        raise ValueError(f"Unsupported investment transaction type '{row.get('tx_type')}'")

    common = {
        "tx_date": parse_date(row.get("tx_date")),
        "tx_type": tx_type,
    }
    if target_table == "mf_transactions":
        candidate = AssetTransactionCreate.model_validate(
            {
                **common,
                "scheme_code": row.get("symbol") or None,
                "mf_name": row.get("asset_name") or row.get("mf_name") or row.get("symbol"),
                "units": row.get("quantity") or row.get("units"),
                "nav": row.get("price") or row.get("nav"),
            }
        ).model_dump(exclude_none=True, mode="json")
        payload = {key: candidate[key] for key in ("tx_date", "tx_type", "scheme_code", "mf_name", "units", "nav") if key in candidate}
    else:
        candidate = AssetTransactionCreate.model_validate(
            {
                **common,
                "symbol": row.get("symbol"),
                "quantity": row.get("quantity") or row.get("units"),
                "price": row.get("price") or row.get("nav"),
                "fees": row.get("fees") or 0,
            }
        ).model_dump(exclude_none=True, mode="json")
        payload = {key: candidate[key] for key in ("tx_date", "tx_type", "symbol", "quantity", "price", "fees") if key in candidate}
    payload["asset_type"] = {
        "equity_transactions": "Equity",
        "mf_transactions": "Mutual Fund",
        "etf_transactions": "ETF",
        "crypto_transactions": "Crypto",
    }[target_table]
    payload["source"] = str(row.get("platform") or "file_import")[:120]
    payload["_target_table"] = target_table
    return payload


def parse_financial_file(filename: str, content: bytes) -> tuple[str, list[dict[str, Any]]]:
    parsed = analyze_financial_file(filename, content)
    return parsed.target, parsed.rows


def analyze_financial_file(
    filename: str,
    content: bytes,
    target_override: str | None = None,
    mapping_override: dict[str, str] | None = None,
) -> ParsedFinancialFile:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        rows, mapping = parse_csv_document(content, mapping_override)
    elif extension == ".xlsx":
        rows, mapping = parse_xlsx_document(content, mapping_override)
    elif extension == ".pdf":
        rows = parse_cas_pdf(content)
        target = "mutual_fund_holdings"
        recognized = sorted(set().union(*(row.keys() for row in rows))) if rows else []
        return ParsedFinancialFile(target, rows, {}, recognized, [], "cas_parser")
    else:
        raise ValueError("Supported financial import formats are CSV, XLSX, and text-based CAS PDF")
    if not rows:
        raise ValueError("The file contains no data rows")
    if target_override:
        if target_override not in TARGET_REQUIREMENTS:
            raise ValueError("The selected financial record type is not supported")
        target = target_override
    else:
        target = detect_target(set(rows[0]))
    accepted = _accepted_columns(target)
    normalized_headers = set(mapping.values())
    recognized = sorted(normalized_headers & accepted)
    ignored = sorted(source for source, canonical in mapping.items() if canonical not in accepted)
    method = "user_override" if mapping_override else "deterministic"
    return ParsedFinancialFile(target, rows, mapping, recognized, ignored, method)


def mapping_diagnostics(filename: str, content: bytes) -> dict[str, Any] | None:
    """Describe a tabular file so the UI can offer a safe mapping fallback."""
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        rows, mapping = parse_csv_document(content)
    elif extension == ".xlsx":
        rows, mapping = parse_xlsx_document(content)
    else:
        return None
    if not rows:
        return None

    normalized = set(mapping.values())
    scores = {
        target: len(normalized & _accepted_columns(target))
        for target in TARGET_REQUIREMENTS
    }
    suggested_target = max(scores, key=scores.get) if max(scores.values(), default=0) else None
    targets = {
        target: {
            "required_columns": required,
            "available_columns": sorted(_accepted_columns(target)),
        }
        for target, required in TARGET_REQUIREMENTS.items()
    }
    return {
        "code": "mapping_required",
        "message": "Automatic recognition could not safely match this file to a supported financial record. You can use advanced column mapping without editing the original file.",
        "source_columns": list(mapping),
        "automatic_mapping": mapping,
        "suggested_target": suggested_target,
        "targets": targets,
        "clarification_options": ["holdings_snapshot", "investment_history", "bank_transactions"],
    }


def _accepted_columns(target: str) -> set[str]:
    model_columns = {
        "equity_holdings": set(EquityCreate.model_fields) | {"current_price"},
        "mutual_fund_holdings": set(MutualFundCreate.model_fields) | {"current_nav"},
        "etf_holdings": set(ETFCreate.model_fields) | {"current_price"},
        "crypto_holdings": set(CryptoCreate.model_fields) | {"current_price"},
        "bank_transactions": {
            "account_id", "account", "tx_date", "description", "amount", "tx_type",
            "type", "category", "debit", "credit",
        },
        "mixed_asset_transactions": {
            "transaction_id", "asset_type", "tx_date", "tx_type", "symbol", "asset_name",
            "quantity", "units", "price", "nav", "amount", "platform", "fees",
        },
        "equity_transactions": {"tx_date", "tx_type", "symbol", "quantity", "price", "fees", "platform"},
        "mf_transactions": {"tx_date", "tx_type", "scheme_code", "mf_name", "units", "nav", "platform"},
        "etf_transactions": {"tx_date", "tx_type", "symbol", "quantity", "price", "fees", "platform"},
        "crypto_transactions": {"tx_date", "tx_type", "symbol", "quantity", "price", "fees", "platform"},
    }
    return model_columns[target]


def tabular_column_profiles(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Build privacy-safe semantic evidence without returning raw financial rows."""
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        rows, mapping = parse_csv_document(content)
    elif extension == ".xlsx":
        rows, mapping = parse_xlsx_document(content)
    else:
        return []
    total = max(len(rows), 1)
    profiles: list[dict[str, Any]] = []
    category_hint = re.compile(r"(^|_)(type|side|action|activity|direction|nature|mode|class|category|asset|instrument|product|transaction)(_|$)")
    for source, canonical in mapping.items():
        values = [str(row.get(canonical) or "").strip() for row in rows]
        populated = [value for value in values if value]
        numeric = [_looks_numeric(value) for value in populated]
        dates = [_looks_like_date(value) for value in populated]
        distinct = sorted(set(populated))
        profile: dict[str, Any] = {
            "source_column": source[:120],
            "normalized_header": canonical,
            "non_empty_rate": round(len(populated) / total, 3),
            "numeric_rate": round(sum(numeric) / max(len(populated), 1), 3),
            "date_rate": round(sum(dates) / max(len(populated), 1), 3),
            "distinct_count": len(distinct),
            "value_shape_examples": list(dict.fromkeys(_value_shape(value) for value in populated[:20]))[:4],
        }
        if category_hint.search(normalize_header(source)) and 0 < len(distinct) <= 20:
            profile["category_values"] = [value[:50] for value in distinct]
        profiles.append(profile)
    return profiles


def _looks_numeric(value: str) -> bool:
    try:
        float(value.replace(",", "").replace("₹", "").replace("$", ""))
        return True
    except ValueError:
        return False


def _looks_like_date(value: str) -> bool:
    candidate = value.strip()
    for parser in (
        lambda: datetime.fromisoformat(candidate.replace("Z", "+00:00")),
        lambda: datetime.strptime(candidate, "%d/%m/%Y"),
        lambda: datetime.strptime(candidate, "%d-%m-%Y"),
        lambda: datetime.strptime(candidate, "%m/%d/%Y"),
    ):
        try:
            parser()
            return True
        except ValueError:
            continue
    return False


def _value_shape(value: str) -> str:
    if _looks_like_date(value):
        return "DATE"
    if _looks_numeric(value):
        return "DECIMAL" if "." in value else "INTEGER"
    shaped = re.sub(r"[A-Za-z]", "A", value)
    shaped = re.sub(r"\d", "9", shaped)
    return re.sub(r"(.)\1{3,}", r"\1\1\1", shaped)[:40]
